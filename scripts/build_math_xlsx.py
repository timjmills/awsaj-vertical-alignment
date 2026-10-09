"""Build the extended K-12 CC + EE math standards workbook.

Keeps the school's K-5 tabs exactly as they are (copied from CCEE Math Standards 2024 V1 Complete.xlsx)
and adds Grade 6-8 and high-school course tabs in the same layout, with draft "I can" descriptors.

Usage: python3 build_math_xlsx.py <out.xlsx>
"""
import json, sys, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = sys.argv[1]
wb = openpyxl.load_workbook("data/source/CCEE_Math_Standards_2024_V1.xlsx")
if "Elementary" in wb.sheetnames and wb["Elementary"].max_row <= 1:
    del wb["Elementary"]

rows = json.load(open("data/math_6_12_merged.json"))
desc = {}
for f in sorted(glob.glob("data/descriptors/math_desc_[0-9].json")):
    desc.update(json.load(open(f)))
desc.update(json.load(open("data/descriptors/math_desc_recheck.json")))
missing = [r["code"] for r in rows if r["code"] not in desc]
assert not missing, missing

BLUE = PatternFill("solid", fgColor="FFC9DAF8")
GREEN = PatternFill("solid", fgColor="FFBAD682")
MINT = PatternFill("solid", fgColor="FFDFFFEB")
GREY = PatternFill("solid", fgColor="FFF3F3F3")
thin = Side(style="thin", color="FF999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F12 = Font(name="Calibri", size=12)
F12B = Font(name="Calibri", size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

HEAD = ["CC Cluster Statement", "Quarter Taught", "Power Standards", "CC Code", "CC Code & Standard",
        "CC Descriptor", "EE Code", "EE Code & Standard", "EE Level Descriptor", "Assessment Links",
        "WI Tag", "Notes"]
WIDTHS = [22, 14.75, 14, 16, 54, 61, 16, 46.88, 60, 26.88, 12, 30]

TABS = [("G6", "Grade 6 Math"), ("G7", "Grade 7 Math"), ("G8", "Grade 8 Math"),
        ("Algebra 1", "Algebra 1"), ("Geometry", "Geometry"), ("Algebra 2", "Algebra 2"),
        ("Pre-Calculus", "Pre-Calculus")]


def cc_desc(d):
    return (f"CC Advanced:\n{d['cc_advanced']}\n\nCC At Target:\n{d['cc_at_target']}\n\n"
            f"CC Approaching:\n{d['cc_approaching']}\n\nCC Emerging:\n{d['cc_emerging']}")


def ee_desc(d, r):
    t = d["ee_at_target"]
    if not r["ee_code"] or t.startswith("Not applicable"):
        return t
    return f"EE At Target:\n{t}"


for tab, course in TABS:
    ws = wb.create_sheet(tab)
    for i, h in enumerate(HEAD, 1):
        c = ws.cell(1, i, h)
        c.font = F12B; c.alignment = WRAP; c.border = BOX
        c.fill = BLUE if i <= 6 else GREEN if i <= 10 else GREY
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    rr = 2
    sub = [r for r in rows if r["course"] == course]
    domain, cluster_start, cluster_key = None, None, None

    def close_cluster(end):
        if cluster_start is not None and end > cluster_start:
            ws.merge_cells(start_row=cluster_start, start_column=1, end_row=end, end_column=1)

    for r in sub:
        if r["domain"] != domain:
            close_cluster(rr - 1); cluster_start = cluster_key = None
            domain = r["domain"]
            c = ws.cell(rr, 1, f"{domain} ({r['domain_code']})")
            c.font = Font(name="Calibri", size=20, bold=True); c.fill = MINT
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=len(HEAD))
            rr += 1
        key = r["code"].rsplit(".", 1)[0]
        if key != cluster_key:
            close_cluster(rr - 1)
            cluster_key, cluster_start = key, rr
            c = ws.cell(rr, 1, r["cluster"] + (" (M)" if r["major_cluster"] else ""))
            c.font = Font(name="Calibri", size=14, bold=True); c.fill = MINT; c.alignment = WRAP
        d = desc[r["code"]]
        ee_std = f"{r['ee_code']} {r['ee_text']}".strip() if r["ee_code"] else r["ee_text"]
        vals = [None, None, None, r["code"], f"{r['code']} {r['cc_text']}", cc_desc(d),
                r["ee_code"], ee_std, ee_desc(d, r), None, r["wi_tags"].replace("+", "(+)"), None]
        for i, v in enumerate(vals, 1):
            if i == 1:
                continue
            c = ws.cell(rr, i, v)
            c.font = F12; c.alignment = WRAP_TOP if i in (5, 6, 8, 9) else WRAP; c.border = BOX
        ws.cell(rr, 1).border = BOX
        rr += 1
    close_cluster(rr - 1)

# About tab at the front
ab = wb.create_sheet("About", 0)
notes = [
    ("Awsaj Academy K-12 Math Standards (Wisconsin CC + EE)", True),
    ("", False),
    ("What this workbook is", True),
    ("KG-G5 tabs: copied unchanged from 'CCEE Math Standards 2024 V1 Complete.xlsx' (the school's original stays untouched).", False),
    ("G6, G7, G8 and the high-school course tabs: new. Same column layout as the K-5 tabs.", False),
    ("", False),
    ("Sources", True),
    ("CC standards: Wisconsin Standards for Mathematics (DPI, May 2021).", False),
    ("EE standards: Wisconsin Essential Elements for Mathematics (DPI, 2022).", False),
    ("", False),
    ("Please review", True),
    ("All 'I can' descriptors on the G6 to Pre-Calculus tabs are DRAFTS written for committee review (4 CC levels + EE At Target).", False),
    ("High-school course placement is a suggestion based on the traditional Algebra 1, Geometry, Algebra 2 pathway. Wisconsin tags: F2Y = first two years of high school; (+) = advanced, beyond F2Y.", False),
    ("Suggested grade for each course: Algebra 1 = Grade 9, Geometry = Grade 10, Algebra 2 = Grade 11, Pre-Calculus = Grade 12 (matches the 2026-27 High School folders).", False),
    ("(M) after a cluster = a major cluster in the Wisconsin standards.", False),
    ("Quarter Taught, Power Standards and Assessment Links are left blank for teams to fill in.", False),
]
for i, (t, b) in enumerate(notes, 1):
    c = ab.cell(i, 1, t)
    c.font = Font(name="Calibri", size=16 if i == 1 else 12, bold=b)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ab.column_dimensions["A"].width = 120
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
