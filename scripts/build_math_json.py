"""Write data/standards/math.json in the shared schema (docs/standards_team_brief.md).

K-5 rows come from the school's CCEE Math sheet (descriptors copied); 6-12 from the parsed Wisconsin docs + drafts.
Usage: python3 build_math_json.py
"""
import json, re, glob
import openpyxl

SRC = "data/source/CCEE_Math_Standards_2024_V1.xlsx"
wb = openpyxl.load_workbook(SRC)
LEVELS = [("cc_advanced", r"CC\s*Advanced"), ("cc_at_target", r"CC\s*(At|On)\s*Target"),
          ("cc_approaching", r"CC\s*Approaching"), ("cc_emerging", r"CC\s*Emerging")]


def split_levels(cell):
    out = {k: "" for k, _ in LEVELS}
    if not cell:
        return out
    t = str(cell)
    marks = []
    for k, pat in LEVELS:
        m = re.search(pat + r"\s*:?", t, re.I)
        if m:
            marks.append((m.start(), m.end(), k))
    marks.sort()
    for i, (s, e, k) in enumerate(marks):
        nxt = marks[i + 1][0] if i + 1 < len(marks) else len(t)
        out[k] = t[e:nxt].strip()
    return out


def ee_target(cell):
    if not cell:
        return ""
    t = str(cell).strip()
    t = re.sub(r"^EE\s*(At|On)\s*Target\s*:?", "", t, flags=re.I).strip()
    return t


def clean(t):
    t = re.sub(r"\s+", " ", str(t or "")).strip()
    return t.replace("—", ", ").replace("–", "-")


rows = []
GRADE = {"KG": "K", "G1": "1", "G2": "2", "G3": "3", "G4": "4", "G5": "5"}
for tab, g in GRADE.items():
    ws = wb[tab]
    merged, banners = {}, set()
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1 and rng.max_col >= 8:
            banners.add(rng.min_row)
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = v
    val = lambda r, c: merged.get((r, c), ws.cell(r, c).value)
    strand = ""
    for r in range(2, ws.max_row + 1):
        code = val(r, 4)
        a = val(r, 1)
        if r in banners or (not code and a and all(not val(r, c) for c in range(4, 9))):
            strand = clean(a); continue
        if not code or not str(code).strip().startswith("M."):
            continue
        code = str(code).strip()
        text = clean(val(r, 5))
        text = re.sub(r"^" + re.escape(code) + r"\.?\s*", "", text)
        ee_code = clean(val(r, 7))
        ee_full = clean(val(r, 8))
        ee_text = re.sub(r"^" + re.escape(ee_code) + r"\.?\s*", "", ee_full) if ee_code else ee_full
        d = split_levels(val(r, 6))
        d["ee_at_target"] = clean(ee_target(val(r, 9)))
        d = {k: clean(v) for k, v in d.items()}
        rows.append({
            "subject": "Math", "framework": "WI-CC-MATH", "code": code, "grade": g, "grade_band": "",
            "grade_is_suggested": False, "division": "Elementary", "course": "",
            "strand": strand, "cluster": clean(val(r, 1)), "text": text,
            "ee_code": ee_code, "ee_text": ee_text, "crosswalk": [], "descriptors": d,
            "descriptor_source": "school",
            "power_standard": (True if val(r, 3) else None),
            "source": "CCEE Math Standards 2024 V1 Complete.xlsx (school)", "text_source": "school sheet",
        })

# one CC standard can sit beside several EE rows (merged cells): combine into one row
combined, order = {}, []
for r in rows:
    k = (r["code"], r["grade"])
    if k not in combined:
        combined[k] = r; order.append(k); continue
    c = combined[k]
    if r["ee_code"] and r["ee_code"] not in c["ee_code"]:
        c["ee_code"] = "; ".join(x for x in [c["ee_code"], r["ee_code"]] if x)
        c["ee_text"] = " | ".join(x for x in [c["ee_text"], r["ee_text"]] if x)
        c["descriptors"]["ee_at_target"] = " | ".join(x for x in [c["descriptors"]["ee_at_target"], r["descriptors"]["ee_at_target"]] if x)
rows = [combined[k] for k in order]

# ---- fixes from the independent check (Oct 9) ----
CODE_FIX = {"M.1.OA.C.5.a.b": "M.1.OA.C.5.b"}
fixed = []
for r in rows:
    r["code"] = CODE_FIX.get(r["code"], r["code"])
    if "[WI.2010" in r["code"]:
        r["code"] = r["code"].split()[0]
    r["text"] = re.sub(r"^M\.\S+\s*\[WI\.2010\.\s*\S+\]\.?\s*", "", r["text"])
    r["text"] = re.sub(r"^" + re.escape(r["code"].rsplit(".", 1)[0]) + r"\s+", "", r["text"]) if r["code"].endswith(".a") else r["text"]
    # EE text "Not applicable." whose "See ..." landed in the descriptor column
    d = r["descriptors"]
    if r["ee_text"].lower().startswith("not applicable"):
        if d["ee_at_target"].startswith("See"):
            r["ee_text"] = (r["ee_text"].rstrip(".") + ". " + d["ee_at_target"]).strip()
        d["ee_at_target"] = r["ee_text"]
    # split EE levels run together in the school cell
    t = d["ee_at_target"]
    for key, pat in [("ee_emerging", r"EE\s*Emerging\s*:?"), ("ee_approaching", r"EE\s*Approaching\s*:?"), ("ee_advanced", r"EE\s*Advanced\s*:?")]:
        m = re.search(pat, t, re.I)
        if m:
            nxt = [m2.start() for m2 in re.finditer(r"EE\s*(Emerging|Approaching|Advanced|At Target|On Target)\s*:?", t, re.I) if m2.start() > m.start()]
            d[key] = t[m.end():(nxt[0] if nxt else len(t))].strip()
    cut = re.search(r"EE\s*(Emerging|Approaching|Advanced)\s*:?", t, re.I)
    if cut:
        d["ee_at_target"] = t[:cut.start()].strip()
    # 1.MD.B.3 has no CC parts: one row paired with all EE parts
    if r["code"] in ("M.1.MD.B.3.b", "M.1.MD.B.3.c", "M.1.MD.B.3.d") and r["text"].lower().startswith("not applicable"):
        base = next(x for x in fixed if x["code"] == "M.1.MD.B.3")
        base["ee_code"] += "; " + r["ee_code"]
        base["ee_text"] += " | " + r["ee_text"]
        base["descriptors"]["ee_at_target"] += " | " + d["ee_at_target"]
        continue
    if r["code"] == "M.1.MD.B.3.a":
        r["code"] = "M.1.MD.B.3"
        r["text"] = re.sub(r"^M\.1\.MD\.B\.3\s*", "", r["text"])
    r["descriptor_flag"], r["descriptors_suggested"] = "", None
    fixed.append(r)
rows = fixed
for r in rows:
    if r["code"] == "M.1.OA.B.4":
        r["descriptor_flag"] = "School descriptor may not match this standard"
        r["descriptor_flag_note"] = "School CC Emerging describes properties of operations (1.OA.B.3), not subtraction as an unknown-addend problem."
        r["descriptors_suggested"] = dict(r["descriptors"], cc_emerging="I can use objects to find how many more I need to add to one number to make another number within 10, with help.")

merged6 = json.load(open("data/math_6_12_merged.json"))
desc = {}
for f in sorted(glob.glob("data/descriptors/math_desc_[0-9].json")):
    desc.update(json.load(open(f)))
desc.update(json.load(open("data/descriptors/math_desc_recheck.json")))
for r in merged6:
    g = r["grade"]
    rows.append({
        "subject": "Math", "framework": "WI-CC-MATH", "code": r["code"], "grade": g, "grade_band": "HS" if r["course_is_suggested"] else "",
        "grade_is_suggested": r["course_is_suggested"], "division": "Middle" if int(g) <= 8 else "High",
        "course": "" if not r["course_is_suggested"] else r["course"],
        "strand": r["domain"], "cluster": r["cluster"], "text": clean(r["cc_text"]),
        "ee_code": r["ee_code"], "ee_text": clean(r["ee_text"]), "crosswalk": [],
        "descriptors": {k: clean(v) for k, v in desc[r["code"]].items()}, "descriptor_source": "draft",
        "descriptor_flag": "", "descriptors_suggested": None,
        "power_standard": None,
        "source": "Wisconsin Standards for Mathematics (2021); Wisconsin Essential Elements for Mathematics (2022)",
        "text_source": "official", "wi_tags": r["wi_tags"], "major_cluster": r["major_cluster"],
    })
json.dump(rows, open("data/standards/math.json", "w"), indent=1, ensure_ascii=False)
from collections import Counter
print(len(rows), sorted(Counter(r["grade"] for r in rows).items(), key=lambda x: ("K0123456789".index(x[0][0]) if x[0] != "10" and x[0] != "11" and x[0] != "12" else 20 + int(x[0]))))
