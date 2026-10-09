"""Build outputs/Social Studies Standards K-12 (Extended).xlsx from the school workbook + social_studies.json.

python3 -I /home/claude/awsaj-va/scripts/social_studies/build_ss_xlsx.py
"""
import collections
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = "/home/claude/awsaj-va"
sys.path.insert(0, f"{ROOT}/scripts/social_studies")
from school_flags import FLAGS, MASTER_DIFFS  # noqa: E402

SRC_XLSX = f"{ROOT}/data/source/social_studies/Social Studies Standards.xlsx"
OUT = f"{ROOT}/outputs/Social Studies Standards K-12 (Extended).xlsx"
BLUE, GREEN, LINKS, BANNER, SUB = "C9DAF8", "BAD682", "CFE2F3", "D9D9D9", "F3F3F3"
thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
GRADES = ["K", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"]
STRAND_ORDER = ["Inquiry Practices and Processes", "Behavioral Sciences", "Economics", "Geography", "History", "Political Science"]

rows = json.load(open(f"{ROOT}/data/standards/social_studies.json"))
meta = json.load(open(f"{ROOT}/data/source/social_studies/build_meta.json"))


def perf(d):
    return (f"Emerging: {d['cc_emerging']}\n\nApproaching: {d['cc_approaching']}\n\n"
            f"At Target: {d['cc_at_target']}\n\nAdvanced: {d['cc_advanced']}")


def fill(c):
    return PatternFill("solid", fgColor=c)


def header(ws, cols):
    for i, (name, width, color) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name="Calibri", bold=True, size=12)
        c.fill = fill(color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def banner(ws, r, ncol, text, color=BANNER, size=12):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(name="Calibri", bold=True, size=size)
    c.fill = fill(color)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 32 if len(text) > 110 else 20


def put(ws, r, vals):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name="Calibri", size=11)
        c.alignment = WRAP
        c.border = BORDER


def wi_tab(wb, title, sel, k5=False):
    ws = wb.create_sheet(title)
    cols = []
    if k5:
        cols.append(("Grade", 7, BLUE))
    cols += [("WI-SS Code", 16, BLUE), ("Learning Priority", 22, BLUE), (f"{title} WI-SS Standard (performance indicator)", 55, BLUE),
             ("Grade Band", 9, BLUE)]
    if k5:
        cols.append(("AERO-SS Crosswalk (suggested)", 16, BLUE))
    cols += [("Power Standard", 10, BLUE), ("Performance Level (DRAFT 'I can')", 70, BLUE), ("Assessment Links", 18, LINKS),
             ("EE Code", 13, GREEN), ("Essential Element (WI EE 2022)", 40, GREEN), ("EE At Target (DRAFT 'I can')", 36, GREEN)]
    header(ws, cols)
    n = len(cols)
    r = 2
    key = lambda x: (GRADES.index(x["grade"]), STRAND_ORDER.index(x["strand"]), x["cluster"].split(" | ")[0], x["code"])
    last_g = last_s = last_std = None
    for x in sorted(sel, key=key):
        std = x["cluster"].split(" | ")[0]
        if k5 and x["grade"] != last_g:
            banner(ws, r, n, f"Grade {x['grade']}" if x["grade"] != "K" else "Kindergarten", color=BLUE, size=13)
            r += 1
            last_s = last_std = None
        if x["strand"] != last_s:
            banner(ws, r, n, x["strand"].upper())
            r += 1
            last_std = None
        if std != last_std:
            banner(ws, r, n, std, color=SUB, size=11)
            r += 1
        last_g, last_s, last_std = x["grade"], x["strand"], std
        lp = x["cluster"].split(" | ")[1]
        vals = []
        if k5:
            vals.append(x["grade"])
        vals += [x["code"], lp, x["text"], x["grade_band"] + (" (course placement suggested)" if x["grade_is_suggested"] else "")]
        if k5:
            vals.append(", ".join(x["crosswalk"]))
        vals += ["", perf(x["descriptors"]), "", x["ee_code"], x["ee_text"], x["descriptors"]["ee_at_target"]]
        put(ws, r, vals)
        r += 1
    return ws


def aero_tab(wb, sel):
    ws = wb.create_sheet("AERO-SS K-5 (Full)")
    cols = [("Grade", 7, BLUE), ("AERO Code", 10, BLUE), ("AERO Standard (official text, AERO PDF)", 50, BLUE),
            ("School report-card wording (if listed)", 40, BLUE), ("Placement source", 18, BLUE),
            ("Power Standard", 10, BLUE), ("Performance Level", 70, BLUE), ("Descriptor source", 11, BLUE),
            ("WI-SS Crosswalk (suggested)", 22, BLUE), ("Committee flag", 40, LINKS),
            ("Suggested Performance Level (DRAFT, for flagged rows)", 70, BLUE)]
    header(ws, cols)
    n = len(cols)
    r = 2
    last_g = last_s = None
    for x in sel:
        m = meta[f"AERO-SS|{x['code']}|{x['grade']}"]
        if x["grade"] != last_g:
            banner(ws, r, n, f"Grade {x['grade']}" if x["grade"] != "K" else "Kindergarten", color=BLUE, size=13)
            r += 1
            last_s = None
        if x["strand"] != last_s:
            banner(ws, r, n, x["strand"])
            r += 1
        last_g, last_s = x["grade"], x["strand"]
        put(ws, r, [x["grade"], x["code"], x["text"], m.get("_school_printed", ""), m["_placement"], "",
                    perf(x["descriptors"]), x["descriptor_source"], ", ".join(x["crosswalk"]),
                    "; ".join(f for f in (x["descriptor_flag"], FLAGS.get((x["grade"], x["code"]), "")) if f),
                    perf(x["descriptors_suggested"]) if x["descriptors_suggested"] else ""])
        if x["descriptor_flag"]:
            for ci in range(1, n + 1):
                ws.cell(row=r, column=ci).fill = fill("FFF2CC")
        r += 1
    return ws


def about(wb, counts, aero_counts):
    ws = wb.create_sheet("About", 0)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    r = 1

    def line(a, b="", bold=False, color=None):
        nonlocal r
        ca = ws.cell(row=r, column=1, value=a)
        cb = ws.cell(row=r, column=2, value=b)
        for c in (ca, cb):
            c.font = Font(name="Calibri", size=11, bold=bold)
            c.alignment = WRAP
            if color:
                c.fill = fill(color)
        r += 1

    line("Social Studies Standards K-12 (Extended)", "Awsaj Academy vertical alignment draft. Built 9 Oct 2026. Original school workbook in Drive is unchanged; this is an extended copy.", True, BLUE)
    line("")
    line("TABS", "", True, BANNER)
    line("KG, G1, G2, G3, G4, G5", "The school's original AERO report-card tabs, copied UNCHANGED.")
    line("AERO-SS K-5 (Full)", "Every AERO K-5 performance indicator at its AERO Learning Progression grade, plus extra rows where the school sheet lists a code at a different grade. School descriptors are copied where the school sheet has them; other rows have DRAFT descriptors. Includes suggested WI-SS crosswalk and committee flags.")
    line("WI-SS K-5", "Wisconsin Standards for Social Studies (2018) performance indicators K-5, paired with Wisconsin Essential Elements (2022), with a suggested AERO-SS crosswalk column.")
    line("G6, G7, G8", "WI-SS grade band 6-8 ('m') indicators. The band is shown in every grade of the band (6, 7 and 8).")
    line("Geography, World History 10, Economics, World History 12", "WI-SS grade band 9-12 ('h') indicators placed into the school's HS courses (placement is suggested; see below).")
    line("")
    line("SOURCES", "", True, BANNER)
    line("WI-SS", "Wisconsin Standards for Social Studies (2018), Wisconsin DPI: https://dpi.wi.gov/sites/default/files/imce/standards/pdf/2018_WI_Social_Studies_Standards.pdf (text parsed from the official PDF).")
    line("WI EE", "Wisconsin Essential Elements for Social Studies (adopted Nov 2021, 2022 PDF): https://dpi.state.wi.us/sites/default/files/imce/literacy-mathematics/pdf/Essential_Elements_for_Social_Studies_2022.pdf. EEs are written per standard (e.g. SS.EE.Hist1), so every indicator under a standard shares that standard's EE.")
    line("AERO-SS", "School copy 'Social Studies Standards K-5 Aero .pdf' (AERO Social Studies Curriculum Framework, K-5 Learning Progression, pp.10-16). Grade = the Learning Progression column in which the indicator is printed.")
    line("School sheets", "'Social Studies Standards.xlsx' (K-5 report-card tabs with 'I can' performance levels) and 'Social Studies Master Standards.xlsx' (K-5 matrix).")
    line("")
    line("HOW GRADES ARE ASSIGNED", "", True, BANNER)
    line("WI-SS K-8", "WI codes end in a band or a grade: .e = K-2, .i = 3-5, .m = 6-8 (row repeated in every grade of the band, grade_band set). Where the document prints a grade-level code (e.g. SS.BH1.a.2, SS.Econ2.a.3-4, SS.Geog2.a.K-1) it is used exactly: that indicator appears only in the grade(s) named.")
    line("WI-SS 9-12", "The 2018 standards give one 9-12 ('h') indicator per learning priority. Rows are placed into the school's HS courses with grade_is_suggested = true: Geography (Grade 9) = Geography strand; World History 10 (Grade 10) = History strand; Economics (Grade 11) = Economics strand; World History 12 (Grade 12) = History + Political Science + Behavioral Sciences strands; the Inquiry strand is attached to all four courses. History indicators therefore appear in both Grade 10 and Grade 12.")
    line("AERO-SS", "AERO codes are banded (x.2.x = K-2, x.5.x = 3-5; grade_band set on every row). The grade comes from the school report-card sheet where it lists the code (grade_is_suggested = false), otherwise from the AERO Learning Progression column (grade_is_suggested = true). Where the school sheet lists a code in another grade than AERO, an extra row is added in that grade.")
    line("Flagged school descriptors", "Rows shaded light yellow on the AERO-SS K-5 (Full) tab keep the school descriptors (unchanged) but the school descriptor may not match the standard; a DRAFT suggested descriptor for the real standard is in the last column (JSON: descriptor_flag and descriptors_suggested).")
    line("")
    line("ROW COUNTS (JSON)", "", True, BANNER)
    line("WI-SS", "; ".join(f"{g}: {counts[g]}" for g in GRADES) + f". Total {sum(counts.values())} rows from 254 unique indicators (64 learning priorities x 4 bands, minus Econ3.c which has no K-2 or 3-5 indicator).")
    line("AERO-SS", "; ".join(f"{g}: {aero_counts[g]}" for g in GRADES[:6]) + f". Total {sum(aero_counts.values())} rows: 97 AERO K-5 indicators + 10 school-only placements.")
    line("")
    line("WHAT IS DRAFT", "", True, BANNER)
    line("Draft descriptors", "All WI-SS 'I can' descriptors (Emerging, Approaching, At Target, Advanced) and all EE At Target statements are DRAFTS written for this project. Wisconsin and US references in the standards (e.g. 'Wisconsin history', 'founding of the United States') are kept in the official text but localised to Qatar in the draft descriptors. AERO rows without school descriptors also carry draft descriptors (marked 'draft' in the AERO-SS K-5 (Full) tab).")
    line("School descriptors", "Copied exactly from the school sheet (split into the four levels). Not rewritten, including where they do not match the standard (see flags).")
    line("Crosswalk", "AERO-SS to WI-SS links are a best-effort match by meaning, within the same band (AERO x.2.x to WI K-2; AERO x.5.x to WI 3-5). All links are SUGGESTED and need committee review.")
    line("Power standards", "The school sheet's Power Standard column is empty in every tab, so no power standards are marked (null in JSON).")
    line("")
    line("COMMITTEE DECISIONS NEEDED", "", True, BANNER)
    line("1. School wording and descriptor mismatches", "The report-card sheet uses some AERO codes with wording or descriptors that belong to a different indicator. The JSON keeps the official AERO text for each code and copies the school descriptors as they are. Decide whether to re-code these rows or rewrite the descriptors:")
    for (g, c), note in FLAGS.items():
        line(f"   Grade {g}  {c}", note)
    line("2. Master matrix vs AERO PDF", "The 'Social Studies Master Standards.xlsx' matrix places these indicators in a different grade from the AERO Learning Progression (code: Master grade / AERO grade). The JSON follows the AERO PDF; Master placements were not added as rows: " +
         "; ".join(f"{c}: {m} / {a}" for c, m, a in MASTER_DIFFS) + ".")
    line("3. HS course placement", "Confirm the 9-12 strand-to-course placement above, especially whether History indicators belong in both World History 10 and 12, and whether Political Science and Behavioral Sciences belong in World History 12.")
    line("4. US/Wisconsin-specific content", "Some WI-SS indicators name Wisconsin, US founding documents or tribal nations (e.g. SS.PS1.b, SS.PS2.a.i, SS.Geog1.c.4-5). Decide whether to adapt these to Qatar for local use.")
    line("5. Crosswalk", "Review the suggested AERO to WI-SS links before the site shows them as 'same standard, other framework'.")
    line("")
    line("TEXT NOTES", "", True, BANNER)
    line("AERO text fixes", "Characters lost when extracting the AERO PDF were restored from the page images: 4.2.e ('act in one's own culture'), 5.5.h ('acceptance' and 'tolerance'), 8.2.a (quotation marks), and the line-break hyphens in 3.5.d 'human-made' and 7.5.a 'non-renewable'.")
    line("WI text", "Line-break hyphens rejoined; learning priority labels normalised (e.g. 'Geog2d.' to 'Geog2.d:', 'Econ 3.c:' to 'Econ3.c:'). The en dash in SS.Econ1.a.h 'cost-benefit' is written as a hyphen (project style rule). The EE PDF prints 'SS:.Econ1' (typo), read as SS.Econ1.")
    return ws


def main():
    wb = openpyxl.load_workbook(SRC_XLSX)
    wi = [x for x in rows if x["framework"] == "WI-SS"]
    aero = [x for x in rows if x["framework"] == "AERO-SS"]
    counts = collections.Counter(x["grade"] for x in wi)
    aero_counts = collections.Counter(x["grade"] for x in aero)
    aero_tab(wb, aero)
    wi_tab(wb, "WI-SS K-5", [x for x in wi if x["division"] == "Elementary"], k5=True)
    for g in ("6", "7", "8"):
        wi_tab(wb, f"G{g}", [x for x in wi if x["grade"] == g])
    for g, name in (("9", "Geography"), ("10", "World History 10"), ("11", "Economics"), ("12", "World History 12")):
        wi_tab(wb, name, [x for x in wi if x["grade"] == g])
    about(wb, counts, aero_counts)
    wb.active = 0
    wb.save(OUT)
    print("saved", OUT, wb.sheetnames)


if __name__ == "__main__":
    main()
