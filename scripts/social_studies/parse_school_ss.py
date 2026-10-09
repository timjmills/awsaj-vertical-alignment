"""Parse the school's K-5 Social Studies report-card workbook (AERO codes + 'I can' performance levels).

Usage: python3 -I parse_school_ss.py <Social Studies Standards.xlsx> <out.json>
"""
import json
import re
import sys

import openpyxl

SHEETS = {"KG": ("K", "A", "C", "D"), "G1": ("1", "A", "C", "D"), "G2": ("2", "A", "C", "D"),
          "G3": ("3", "A", "C", "D"), "G4": ("4", "C", "D", "E"), "G5": ("5", "A", "C", "D")}
CODE_RE = re.compile(r"^\s*(\d)\.(\d)\.?\s?([a-j])\b\s*[:.]?\s*")
LEVEL_RE = re.compile(r"(?:^|\n)\s*(Emerging|Approaching|At Target|Advanced)\s*:?[ \t]*", re.I)
KEYS = {"emerging": "cc_emerging", "approaching": "cc_approaching", "at target": "cc_at_target", "advanced": "cc_advanced"}


def clean(s):
    s = s.replace("\u00a0", " ")
    s = re.sub(r"\n\s*\d:\s*$", "", s.strip())  # stray "2:" at the end of one G2 cell
    s = re.sub(r"\s*\n\s*", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def parse_levels(txt):
    out = {}
    ms = list(LEVEL_RE.finditer(txt))
    for i, m in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(txt)
        out[KEYS[m.group(1).lower()]] = clean(txt[m.end():end])
    return out


def main(path, out_path):
    wb = openpyxl.load_workbook(path)
    rows = []
    for sheet, (grade, scol, dcol, lcol) in SHEETS.items():
        ws = wb[sheet]
        for r in range(2, ws.max_row + 1):
            sv = ws[f"{scol}{r}"].value
            if not sv or not str(sv).strip():
                continue
            raw = str(sv).strip()
            m = CODE_RE.match(raw)
            if not m:
                print("WARN no code", sheet, r, raw[:50], file=sys.stderr)
                continue
            code = f"{m.group(1)}.{m.group(2)}.{m.group(3)}"
            school_text = clean(raw[m.end():])
            dv = ws[f"{dcol}{r}"].value
            levels = parse_levels(str(dv)) if dv and str(dv).strip() else {}
            pv = ws[f"B{r}"].value if sheet != "G4" else ws[f"B{r}"].value
            rows.append({"sheet": sheet, "row": r, "grade": grade, "code": code, "printed": clean(raw),
                         "school_text": school_text, "descriptors": levels,
                         "power_cell": pv, "links": ws[f"{lcol}{r}"].value})
    json.dump(rows, open(out_path, "w"), indent=1, ensure_ascii=False)
    print(len(rows), "school rows")
    for x in rows:
        print(x["grade"], x["code"], len(x["descriptors"]), "|", x["school_text"][:70], "| power:", x["power_cell"])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
