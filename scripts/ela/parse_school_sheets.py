"""Parse the school's ELA sheets into row dicts.

K-5: data/source/EE ELA New Standards 2025 V1 Complete.xlsx (tabs KG, G1-G5)
9-10: data/source/ela/g910_html/G9-10.html (HTML export of Drive sheet "G9-10 ELA New Standards 2024")
Usage: python3 -I parse_school_sheets.py <k5.xlsx> <g910.html> <out_k5.json> <out_910.json>
"""
import json, re, sys
import openpyxl
from bs4 import BeautifulSoup

PREFIX_STRAND = {"RF": "Reading Foundational Skills", "R": "Reading", "W": "Writing",
                 "SL": "Speaking and Listening", "L": "Language"}
LEVEL = re.compile(r"^\s*(?:CC |EE )*(Advanced|At Target|On Target|Approaching|Emerging)"
                   r"(?: Students will| \(matches the standard\))?\s*:", re.M)


def clean(s):
    if s is None:
        return ""
    s = str(s).replace(" ", " ").replace("—", ", ").replace("–", "-")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return s.strip()


def flat(s):
    return re.sub(r"\s+", " ", clean(s)).strip()


def norm_code(raw, ee=False):
    """'RF2.1' -> 'RF.2.1', ' RF.1.2.b' -> 'RF.1.2.b', 'W.2.4.' -> 'W.2.4', 'W.9.10.9' -> 'W.9-10.9',
    'EERF.K.1.b' -> 'EE.RF.K.1.b'."""
    c = re.sub(r"\s+", "", str(raw or "")).rstrip(".")
    c = c.replace("9.10", "9-10").replace("11.12", "11-12")
    if ee:
        c = re.sub(r"^EE\.?", "", c)
    c = re.sub(r"^(RF|R|W|SL|L)\.?(K|\d+(?:-\d+)?)\.", r"\1.\2.", c)
    return ("EE." + c) if ee and c else c


def strip_code(text, code):
    """Remove the leading code (any spacing/punctuation variant) from a 'Code & Standard' cell."""
    t = flat(text)
    t = re.sub(r"^(EE\.?)?[A-Z]{1,2}\s*\.?\s*(K|\d+(?:[-.]\d+)?)\s*\.\s*\d+(\s*\.\s*[a-z](-[a-z])?)?\s*\.?\s*:?\s*", "", t)
    return t.strip()


def split_levels(raw):
    raw = clean(raw)
    out = {}
    parts = list(LEVEL.finditer(raw))
    if not parts:
        return out, raw
    for i, m in enumerate(parts):
        end = parts[i + 1].start() if i + 1 < len(parts) else len(raw)
        lvl = m.group(1).replace("On Target", "At Target")
        body = raw[m.end():end].strip()
        body = re.sub(r"^(EE\.?)?[A-Z]{1,2}\.(K|\d+(?:-\d+)?)\.\d+\.?[a-z]?\s*:\s*", "", body)
        out[lvl] = (out[lvl] + "\n" + body) if lvl in out else body
    return out, raw[:parts[0].start()].strip()


def descriptors(cc_raw, ee_raw):
    cc, cc_pre = split_levels(cc_raw)
    ee, ee_pre = split_levels(ee_raw)
    d = {"cc_advanced": cc.get("Advanced", ""), "cc_at_target": cc.get("At Target", ""),
         "cc_approaching": cc.get("Approaching", ""), "cc_emerging": cc.get("Emerging", ""),
         "ee_at_target": ee.get("At Target", "")}
    if not cc and cc_pre:      # e.g. "RF.3.3.c. Not Applicable"
        d["cc_at_target"] = strip_code(cc_pre, "")
    if not ee and ee_pre and ee_pre not in ("-",):
        d["ee_at_target"] = strip_code(ee_pre, "")
    # keep the school's extra EE levels so nothing is lost
    for k in ("Advanced", "Approaching", "Emerging"):
        if ee.get(k):
            d["ee_" + k.lower().replace(" ", "_")] = ee[k]
    if ee_pre and ee and ee_pre not in ("-",):
        d["ee_note"] = ee_pre
    if cc_pre and cc:
        d["cc_note"] = cc_pre
    return d


def grade_of(code):
    return code.split(".")[1]


# ---------- K-5 ----------
wb = openpyxl.load_workbook(sys.argv[1])
k5 = []
for tab in ["KG", "G1", "G2", "G3", "G4", "G5"]:
    ws = wb[tab]
    merged_val = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1 and rng.max_col == 1:
            v = ws.cell(rng.min_row, 1).value
            for r in range(rng.min_row, rng.max_row + 1):
                merged_val[r] = v
    concept = ""
    prev_code = None
    for r in range(2, ws.max_row + 1):
        A = merged_val.get(r, ws.cell(r, 1).value)
        C, D, E, F, G, H = (ws.cell(r, k).value for k in range(3, 9))
        if A and not C and not D:
            concept = ""   # strand banner row
            continue
        if not C and not D:
            continue
        if C and not D:    # KG banner row in column C
            continue
        if A:
            concept = flat(A).rstrip(".").strip()
        code_raw = str(C) if C else ""
        if not C:
            dd = flat(D)
            m = re.match(r"^\.\s*([a-z])\.\s+(.*)$", dd, re.S)
            if not m or not prev_code:
                print("skip junk row", tab, r, repr(dd[:40]))
                continue
            code = re.sub(r"\.[a-z]$", "", prev_code) + "." + m.group(1)
            text = m.group(2)
        else:
            code = norm_code(C)
            text = strip_code(D, code)
        prev_code = code
        ee_code = norm_code(F, ee=True) if F and flat(F) not in ("N/A",) else ""
        ee_text = strip_code(G, ee_code) if G else ""
        if not C:
            ee_text = re.sub(r"^\.\s*[a-z]\.\s*", "", flat(G or ""))
        k5.append({
            "tab": tab, "row": r, "code": code, "code_original": flat(code_raw) if flat(code_raw) != code else "",
            "grade": "K" if tab == "KG" else tab[1], "strand": PREFIX_STRAND[code.split(".")[0]],
            "cluster": concept, "text": text, "ee_code": ee_code,
            "ee_code_original": flat(F) if F and flat(F) != ee_code else "",
            "ee_text": ee_text, "descriptors": descriptors(E, H),
            # the Power Standards column is empty on every K-5 tab, so nothing is marked
            "power_standard": True if ws.cell(r, 2).value else None,
        })

# ---------- 9-10 ----------
soup = BeautifulSoup(open(sys.argv[2]).read(), "lxml")
grid, spans = [], {}
for ri, tr in enumerate(soup.find("table").find_all("tr")):
    cells = tr.find_all(["td", "th"])[1:]   # drop the row-number header cell
    row, ci = [], 0
    it = iter(cells)
    out = {}
    col = 0
    pending = list(cells)
    while pending or col in spans:
        if col in spans and spans[col][0] > 0:
            out[col] = spans[col][1]
            spans[col][0] -= 1
            if spans[col][0] == 0:
                del spans[col]
            col += 1
            continue
        if not pending:
            break
        c = pending.pop(0)
        txt = c.get_text("\n")
        rs, cs = int(c.get("rowspan", 1)), int(c.get("colspan", 1))
        for k in range(cs):
            out[col + k] = txt
            if rs > 1:
                spans[col + k] = [rs - 1, txt]
        col += cs
    grid.append([out.get(k, "") for k in range(max(out) + 1 if out else 0)])
g910 = []
for row in grid:
    row = row + [""] * (9 - len(row))
    A, C, D, E, F, G, H = row[0], row[1], row[2], row[3], row[4], row[5], row[6]
    if not flat(C) or not re.match(r"^(RF|R|W|SL|L)\.", flat(C)):
        continue
    code = norm_code(C)
    ee_code = norm_code(F, ee=True) if flat(F) else ""
    g910.append({
        "code": code, "code_original": flat(C) if flat(C) != code else "", "strand": PREFIX_STRAND[code.split(".")[0]],
        "cluster": flat(A), "text": strip_code(D, code), "ee_code": ee_code,
        "ee_code_original": flat(F) if flat(F) and flat(F) != ee_code else "",
        "ee_text": flat(G), "descriptors": descriptors(E, H), "cc_desc_raw": clean(E), "ee_desc_raw": clean(H),
    })

json.dump(k5, open(sys.argv[3], "w"), indent=1, ensure_ascii=False)
json.dump(g910, open(sys.argv[4], "w"), indent=1, ensure_ascii=False)
from collections import Counter
print("K-5 rows", Counter(r["grade"] for r in k5))
print("9-10 rows", len(g910))
