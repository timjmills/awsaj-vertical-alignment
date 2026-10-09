"""Build outputs/ELA Standards K-12 (Extended).xlsx

Starts from a copy of the school's K-5 workbook (KG-G5 tabs kept unchanged), adds G6-G12 tabs in the same
column layout and colours, and an About tab first. Data comes from data/standards/ela.json.
Run from /home/claude/awsaj-va:  python3 -I scripts/ela/build_ela_xlsx.py
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "data/source/EE ELA New Standards 2025 V1 Complete.xlsx"
OUT = "outputs/ELA Standards K-12 (Extended).xlsx"
rows = json.load(open("data/standards/ela.json"))

wb = openpyxl.load_workbook(SRC)
PEACH = PatternFill("solid", fgColor="FFFCE5CD")    # school's CC header colour
GREEN = PatternFill("solid", fgColor="FFD9EAD3")    # school's EE header colour
BANNER = PatternFill("solid", fgColor="FFF9CB9C")   # school's strand banner colour
GREY = PatternFill("solid", fgColor="FFF3F3F3")
DRAFT = PatternFill("solid", fgColor="FFFFF2CC")
thin = Side(style="thin", color="FF000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F12 = Font(name="Calibri", size=12)
F12B = Font(name="Calibri", size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

WIDTHS = {"A": 17.0, "B": 10.13, "C": 10.5, "D": 40, "E": 79.88, "F": 12.5, "G": 40, "H": 70, "I": 30, "J": 34}
STRAND_ORDER = ["Reading", "Writing", "Speaking and Listening", "Language"]

LBL = {"cc_advanced": "CC Advanced", "cc_at_target": "CC At Target", "cc_approaching": "CC Approaching",
       "cc_emerging": "CC Emerging"}


def cc_desc(r):
    d, drafted = r["descriptors"], r["descriptor_draft_fields"]
    parts = []
    for k in ["cc_advanced", "cc_at_target", "cc_approaching", "cc_emerging"]:
        lab = LBL[k] + (" (draft)" if k in drafted and r["descriptor_source"] == "school" else "")
        parts.append(f"{lab}:\n{d.get(k, '')}")
    return "\n\n".join(parts)


def ee_desc(r):
    d, drafted = r["descriptors"], r["descriptor_draft_fields"]
    t = d.get("ee_at_target", "")
    if t.lower().startswith(("not applicable", "n/a")) and not d.get("ee_approaching"):
        return t
    parts = []
    if d.get("ee_advanced"):
        parts.append(f"EE Advanced:\n{d['ee_advanced']}")
    lab = "EE At Target" + (" (draft)" if "ee_at_target" in drafted and r["descriptor_source"] == "school" else "")
    parts.append(f"{lab}:\n{t}")
    for k, lab in (("ee_approaching", "EE Approaching"), ("ee_emerging", "EE Emerging")):
        if d.get(k):
            parts.append(f"{lab}:\n{d[k]}")
    return "\n\n".join(parts)


def note(r):
    bits = []
    if r["grade_band"]:
        bits.append(f"Grade band {r['grade_band']} standard, taught in every grade of the band ({r['course']} tab shown).")
    if r["descriptor_source"] == "draft":
        bits.append("Descriptors: DRAFT for committee review.")
    elif r["descriptor_draft_fields"]:
        bits.append("School descriptors; blank levels filled with drafts: " +
                    ", ".join(LBL.get(k, "EE At Target") for k in r["descriptor_draft_fields"]) + ".")
    if r["text_source"] == "official":
        bits.append("Text: WI Standards for ELA 2020.")
    else:
        bits.append("Text: school G9-10 sheet.")
    if r.get("ee_text_source") == "official":
        bits.append("EE: WI Essential Elements for ELA 2022.")
    if r["code_original"]:
        bits.append(f"Code printed as {r['code_original']} in source.")
    if r["ee_code_original"]:
        bits.append(f"EE code printed as {r['ee_code_original']} in source.")
    if r.get("descriptor_flag"):
        bits.append("FLAG: school descriptor may not match this standard (" + r.get("descriptor_flag_note", "") +
                    ") A suggested draft is in the JSON 'descriptors_suggested'.")
    if r["crosswalk"]:
        bits.append("CCSS-style: " + ", ".join(r["crosswalk"]) + ".")
    return " ".join(bits)


TABS = [("G6", "6"), ("G7", "7"), ("G8", "8"), ("G9", "9"), ("G10", "10"), ("G11", "11"), ("G12", "12")]
for tab, g in TABS:
    sub = [r for r in rows if r["grade"] == g]
    sub.sort(key=lambda r: STRAND_ORDER.index(r["strand"]))   # stable: keeps code order inside strands
    ws = wb.create_sheet(tab)
    band = sub[0]["grade_band"]
    label = f"G{g} CC Code & Standard" + (f" (Grades {band} band)" if band else "")
    head = ["Concept", "Power Standards", "CC Code", label, "CC Descriptor", "EE Code", "EE Code & Standard",
            "EE Descriptor", "Assessment Links", "Source / Notes"]
    for i, h in enumerate(head, 1):
        c = ws.cell(1, i, h)
        c.font = F12B
        c.alignment = CENTER
        c.border = BOX
        c.fill = PEACH if i <= 5 else GREEN if i <= 9 else GREY
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    rr = 2
    strand = None
    cl_start, cl_name = None, None

    def close_cluster(end):
        if cl_start is not None and end > cl_start:
            ws.merge_cells(start_row=cl_start, start_column=1, end_row=end, end_column=1)

    for r in sub:
        if r["strand"] != strand:
            close_cluster(rr - 1)
            cl_start = cl_name = None
            strand = r["strand"]
            c = ws.cell(rr, 1, strand)
            c.font = Font(name="Calibri", size=20, bold=True)
            c.fill = BANNER
            c.alignment = CENTER
            for k in range(1, len(head) + 1):
                ws.cell(rr, k).border = BOX
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=len(head))
            rr += 1
        if r["cluster"] != cl_name:
            close_cluster(rr - 1)
            cl_name, cl_start = r["cluster"], rr
            c = ws.cell(rr, 1, cl_name)
            c.font = Font(name="Calibri", size=14, bold=True)
            c.alignment = CENTER
        ee_std = f"{r['ee_code']} {r['ee_text']}".strip() if r["ee_code"] else r["ee_text"]
        vals = [None, None, r["code"], f"{r['code']} {r['text']}", cc_desc(r), r["ee_code"], ee_std, ee_desc(r),
                None, note(r)]
        for i, v in enumerate(vals, 1):
            c = ws.cell(rr, i)
            if i > 1:
                c.value = v
            c.border = BOX
            if i in (3, 6):
                c.font = F12B
                c.fill = PEACH if i == 3 else GREEN
                c.alignment = CENTER
            elif i > 1:
                c.font = Font(name="Calibri", size=10 if i == 10 else 12)
                c.alignment = WRAP_TOP if i in (4, 5, 7, 8, 10) else WRAP
        if r["descriptor_source"] == "draft" or r["descriptor_draft_fields"] or r.get("descriptor_flag"):
            ws.cell(rr, 10).fill = DRAFT
        rr += 1
    close_cluster(rr - 1)

# ---- About tab first ----
ab = wb.create_sheet("About", 0)
from collections import Counter
cnt = Counter(r["grade"] for r in rows)
lines = [
    ("Awsaj Academy K-12 ELA Standards: Wisconsin ELA (WI-CC-ELA) with Wisconsin / DLM Essential Elements", "title"),
    ("", None),
    ("What this workbook is", "h"),
    ("KG to G5 tabs: copied unchanged from the school's 'EE ELA New Standards 2025 V1 Complete.xlsx'. The original file in Google Drive was not touched.", None),
    ("G6 to G12 tabs: new, in the same column layout and colours as the K-5 tabs. One row per standard (lettered sub-items get their own row, as on the school sheets).", None),
    ("", None),
    ("Sources", "h"),
    ("Grades 6-8 and 11-12 standards text: Wisconsin Standards for English Language Arts (DPI, 2020), dpi.wi.gov/media/42288 (PDF) and media/54346 (Word version used for parsing).", None),
    ("Grades 6-12 Essential Elements: Wisconsin Essential Elements for English Language Arts (DPI, 2022), dpi.wi.gov/media/48585.", None),
    ("Grades 9 and 10: codes, text and descriptors copied from the school's Google Sheet 'G9-10 ELA New Standards 2024'. Where that sheet leaves an EE cell blank, the EE comes from the 2022 DPI document.", None),
    ("", None),
    ("Codes", "h"),
    ("The school's sheets and Wisconsin 2020 use the same codes: R (Reading, literature and informational combined), RF (Reading Foundational Skills, K-5 only), W, SL, L. Example: R.6.1, W.9-10.2.a, SL.11-12.1.d.", None),
    ("Each Wisconsin Reading standard is tagged (RL), (RI) or (RI&RL); the matching CCSS-style codes (e.g. RL.6.1, RI.6.1) are listed in the Notes column and in the JSON 'crosswalk' field.", None),
    ("Grades 9-10 and 11-12 are grade bands: every band standard appears on both grade tabs (G9 and G10; G11 and G12) as English Language Arts 9 / 10 / 11 / 12.", None),
    ("", None),
    ("What is DRAFT (needs committee review)", "h"),
    ("All 'I can' descriptors on the G6, G7, G8, G11 and G12 tabs (4 CC levels + EE At Target) are drafts written for this workbook. The Notes column is shaded yellow on every row that contains draft text.", None),
    ("G9 and G10: descriptors are the school's own; a few blank levels (mostly CC Emerging and EE At Target) were filled with drafts and are labelled '(draft)' in the cell.", None),
    ("", None),
    ("Committee decisions needed", "h"),
    ("1. Confirm that the 9-10 and 11-12 band standards are taught in both grades of each band (as listed here).", None),
    ("2. The school's G9-10 sheet keeps W.9-10.6 as one row; Wisconsin 2020 splits it into W.9-10.6.a-c. G11/G12 follow Wisconsin (W.11-12.6.a-c). Decide which split to use.", None),
    ("3. Wisconsin prints L.7.3 and L.8.3 without lettered sub-items, so they are one row each.", None),
    ("4. Power Standards and Assessment Links are blank on every tab (also blank on the school's K-5 tabs).", None),
    ("5. Literacy standards for history/social studies, science and technical subjects (grades 6-12) are not included here.", None),
    ("6. 46 school descriptor rows (K-5 and G9-10) look like they describe a different standard. They are kept as written but flagged in the JSON with a suggested draft (key 'descriptors_suggested'). On G9/G10 the flag is shown in the Notes column.", None),
    ("", None),
    ("Corrections made in the JSON master list only (the KG-G5 tabs above are the school's file, unchanged)", "h"),
] + [(f"- {x}", None) for x in json.load(open("data/ela/fixlog_round1.json"))] + [
    ("- Blank or placeholder descriptor levels ('N/A', 'Not applicable', '-') on real standards were filled with drafts and listed in 'descriptor_draft_fields'.", None),
    ("- W.4.3.b and W.4.3.c: EE code printed as 'E.W.4.3.b/c' corrected to EE.W.4.3.b/c.", None),
    ("", None),
    ("Rows per grade in the JSON master list (after corrections)", "h"),
    (", ".join(f"{('KG' if g == 'K' else 'G' + g)}: {cnt[g]}" for g in ["K", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"]), None),
]
for i, (t, kind) in enumerate(lines, 1):
    c = ab.cell(i, 1, t)
    c.font = Font(name="Calibri", size=16 if kind == "title" else 12, bold=kind in ("title", "h"))
    c.alignment = Alignment(wrap_text=True, vertical="top")
ab.column_dimensions["A"].width = 130
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
