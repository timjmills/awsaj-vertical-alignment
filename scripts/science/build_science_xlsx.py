"""Build outputs/Science Standards K-12 (Extended).xlsx from the school's 'Science Standards.xlsx' + data/standards/science.json."""
import json, re, shutil, difflib
from collections import OrderedDict, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = '/home/claude/awsaj-va'
SRC = ROOT + '/data/source/science/'
OUT = ROOT + '/outputs/Science Standards K-12 (Extended).xlsx'
shutil.copy(SRC + 'Science Standards.xlsx', OUT)
wb = openpyxl.load_workbook(OUT)
rows = json.load(open(ROOT + '/data/standards/science.json'))
school = json.load(open(SRC + 'school_rows.json'))
notes = json.load(open(SRC + 'build_notes.json'))
ng = json.load(open(SRC + 'ngss_pes.json'))['pes']

BLUE, GREEN, LBLUE, BANNER = 'C9DAF8', 'BAD682', 'CFE2F3', 'DFFFEB'
thin = Side(style='thin', color='000000')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
def fill(c): return PatternFill('solid', fgColor='FF' + c)
WRAP = Alignment(wrap_text=True, vertical='center')
WRAPTOP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')

def header(ws, cols, widths):
    for i, (name, colr) in enumerate(cols, 1):
        c = ws.cell(1, i, name)
        c.font = Font(name='Calibri', size=12, bold=True); c.fill = fill(colr); c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = 'A2'

def banner(ws, r, text, ncols):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(r, 1, text)
    c.font = Font(name='Calibri', size=20, bold=True); c.fill = fill(BANNER); c.alignment = CENTER
    for k in range(1, ncols + 1): ws.cell(r, k).border = BORDER

def perf(d):
    return (f"Emerging:\n{d['cc_emerging']}\n\nApproaching:\n{d['cc_approaching']}\n\n"
            f"At Target:\n{d['cc_at_target']}\n\nAdvanced:\n{d['cc_advanced']}")

DOMAIN_ORDER = ['Life Science', 'Physical Science', 'Earth and Space Science', 'Engineering, Technology, and Applications of Science']
DOMAIN_LABEL = {'Engineering, Technology, and Applications of Science': 'Engineering Design', 'Earth and Space Science': 'Earth Science'}
DISC_ORDER = ['LS1', 'LS2', 'LS3', 'LS4', 'PS1', 'PS2', 'PS3', 'PS4', 'ESS1', 'ESS2', 'ESS3', 'ETS1']
def disc(code): return re.search(r'(PS\d|LS\d|ESS\d|ETS\d)', code).group(1)
def num(code): return int(code.rsplit('-', 1)[1])

def grade_tab(title, grade, course_label):
    ws = wb.create_sheet(title)
    cols = [('Concept', BLUE), (f'{course_label} NGSS Code & Performance Expectation', BLUE), ('Performance Level (DRAFT)', GREEN),
            ('Assessment Links', LBLUE), ('Essential Element (Wisconsin EE for Science)', GREEN), ('EE Performance Level (DRAFT)', GREEN),
            ('Crosswalk: Wisconsin Standards for Science', LBLUE), ('Placement note', LBLUE)]
    header(ws, cols, [34, 60, 95, 30, 60, 45, 34, 40])
    rs = [x for x in rows if x['framework'] == 'NGSS' and x['grade'] == grade]
    rs.sort(key=lambda x: (DISC_ORDER.index(disc(x['code'])), num(x['code'])))
    r = 2
    for dom in DOMAIN_ORDER:
        sub = [x for x in rs if x['strand'] == dom]
        if not sub: continue
        banner(ws, r, DOMAIN_LABEL.get(dom, dom), len(cols)); r += 1
        groups = OrderedDict()
        for x in sub: groups.setdefault(disc(x['code']), []).append(x)
        for dname, items in groups.items():
            start = r
            for x in items:
                concept = x['cluster'].split(' (')[0].split(' ', 1)[1]
                code = x['code']
                if code.startswith('MS-'):
                    pn = f"Grade {grade} (suggested). Basis: {notes['ms_basis'][code]}." if not code.startswith('MS-ETS') else 'MS-ETS1 is placed in all three middle school grades.'
                elif code.startswith('HS-ETS'):
                    pn = 'HS-ETS1 is attached to all four high school courses (Earth Environmental Science, Physics, Biology, Chemistry).'
                else:
                    pn = f"{x['course']} (Grade {grade}), suggested from the HS course model."
                vals = [concept, f"{code} {x['text']}", perf(x['descriptors']), '',
                        (f"{x['ee_code']}\n{x['ee_text']}" if x['ee_code'] else ''), x['descriptors']['ee_at_target'],
                        '\n'.join(c.split(':', 1)[1] for c in x['crosswalk'] if c.startswith('WI-SCI:')), pn]
                for k, v in enumerate(vals, 1):
                    c = ws.cell(r, k, v); c.font = Font(name='Calibri', size=12 if k != 1 else 14, bold=(k == 1))
                    c.alignment = WRAP if k != 3 else WRAPTOP; c.border = BORDER
                r += 1
            if r - 1 > start: ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    return ws

# ---------------------------------------------------------------- About (first)
ab = wb.create_sheet('About', 0)
ab.column_dimensions['A'].width = 150
lines = [
 ('Awsaj Academy K-12 Science Standards (NGSS, AERO Science, Wisconsin Standards for Science, Wisconsin EE)', 'title'),
 ('', ''),
 ('What this workbook is', 'h'),
 ("KG, G1, G2, G3, G4, G5 tabs: copied unchanged from the school's 'Science Standards.xlsx' (K-5 science report-card sheet). The original in Google Drive was not touched.", ''),
 ("'K-5 NGSS-AERO Crosswalk' tab: every K-5 performance expectation in both frameworks side by side, with the Wisconsin crosswalk, EE pairing and review flags.", ''),
 ("G6, G7, G8, Earth Environmental Science, Physics, Biology, Chemistry tabs: new, in the same first four columns as the school tabs (Concept, Code & Standard, Performance Level, Assessment Links) plus EE and crosswalk columns.", ''),
 ("'WI-SCI K-12' tab: the full Wisconsin Standards for Science (crosscutting concepts, practices, disciplinary core ideas, Wisconsin ETS3 sample indicators) with draft descriptors, for reference and crosswalk.", ''),
 ('', ''),
 ('Decision recorded (principal-level user, Tim)', 'h'),
 ('Middle and high school unit plans cite NGSS performance expectation codes (e.g. MS-ESS2-1, HS-PS1-7), so G6-G12 tabs use NGSS.', ''),
 ('Elementary: it is not yet settled whether elementary uses AERO Science or NGSS, so the master list (data/standards/science.json) carries BOTH frameworks for K-5 (framework "AERO-SCI" and "NGSS") and both are selectable on the site.', ''),
 ('AERO Science K-5 performance expectations use the same codes as NGSS. 72 of 78 have identical wording; 6 have small AERO wording differences (marked in the crosswalk tab): K-PS3-2, K-ESS3-1, 1-PS4-2, 1-PS4-3, 4-ESS3-1, 5-ESS1-1. The school report card uses the AERO wording.', ''),
 ('', ''),
 ('Sources', 'h'),
 ('NGSS: "DCI Arrangements of the Next Generation Science Standards" (Achieve, 2013; nextgenscience.org/sites/default/files/AllDCI.pdf). 208 performance expectations, all text official.', ''),
 ('AERO-SCI: "AERO Science PK-5 Framework Progressions Matrix" (school copy, AERO Science PK-5.pdf). K-5 only (78 PEs); the 5 PreK PEs are outside the K-12 grade model and were not included.', ''),
 ('WI-SCI: Wisconsin Standards for Science (DPI, 2017), accessible .docx and printer-friendly PDF from dpi.wi.gov/science/standards. The school sheet "Science Standards CCEE PRIME.xlsx" Sheet1 holds only the 9-12 column, so the full K-12 text was taken from the DPI document; its 9-12 text was checked against Sheet1 and matches.', ''),
 ('EE: Wisconsin Essential Elements for Science, "Science Standards CCEE PRIME.xlsx" Sheet2 (43 EEs: 9 elementary (grade 5), 14 middle school, 20 high school). Each EE names its NGSS PE (e.g. EE.5-PS1-2 pairs with 5-PS1-2). EE text = Target, Precursor and Initial levels as printed.', ''),
 ("K-5 descriptors: copied from the school's 'Science Standards.xlsx' (descriptor_source \"school\"). Level names in the school sheet are Emerging / Approaching / At Target / Advanced and map 1:1 to cc_emerging / cc_approaching / cc_at_target / cc_advanced. 'Descriptor:' labels were dropped; 'Example:' lines were kept.", ''),
 ("'Science Master Standards.xlsx' (Drive): a K-5 matrix that picks one PE per focus area per grade (42 PEs, shortened wording). It adds no new standards; it is shown in the crosswalk tab as 'In K-5 Science Matrix Map' and may be the school's priority list. Committee to decide whether these are power standards (power_standard is left null).", ''),
 ('', ''),
 ('Grade placement of middle and high school PEs (grade_is_suggested = true)', 'h'),
 ('NGSS middle school and high school PEs are grade-banded (6-8, 9-12). Each was placed in one grade/course:', ''),
 ('Middle school, in priority order: (1) Awsaj 2026-27 unit plans in Drive (G6: MS-ETS1-1 to 4, MS-ESS1-1, MS-ESS2-1, MS-ESS2-4, MS-ESS2-5, MS-ESS2-6, MS-ESS3-1, MS-ESS3-2; G8: MS-PS1-2, MS-PS1-5, MS-PS2-1, MS-PS2-2, MS-PS3-2); (2) Awsaj Secondary Science Scope and Sequence (2025) (G6: MS-PS1-1, MS-PS1-4; G7: MS-LS1-6, MS-LS2-3, MS-LS2-4, MS-ESS2-2; G8: MS-ESS1-4, MS-LS4-1, MS-LS4-4, MS-LS4-6, MS-PS2-4, MS-PS3-1); (3) every other PE from NGSS Appendix K, Table 3A (California Integrated Learning Progressions model, the "preferred integrated" course map): Course 1 = Grade 6, Course 2 = Grade 7, Course 3 = Grade 8. MS-ETS1-1 to 4 appear in all three grades.', ''),
 ('Result: G6 has 25 rows, G7 15, G8 27. The current unit plans pull several Earth science PEs into Grade 6 and leave Grade 7 lighter than Appendix K; the committee should confirm.', ''),
 ('High school: Earth Environmental Science = Grade 9 (all HS-ESS); Physics = Grade 10 (HS-PS2, HS-PS3 except PS3-4, HS-PS4); Biology = Grade 11 (all HS-LS); Chemistry = Grade 12 (all HS-PS1 and HS-PS3-4, which is thermal energy transfer and calorimetry). HS-ETS1-1 to 4 are attached to all four courses (one row per course).', ''),
 ('Note: the school scope and sequence also cites some PEs outside these course rules (e.g. Grade 9 cites HS-PS1-2, HS-LS2-1, HS-LS2-3, HS-LS2-5, HS-LS2-7, HS-LS4-1, HS-LS4-6, HS-PS3-3; Biology cites MS-LS1-1 to 4; Chemistry cites HS-PS3-1). Teachers rating Part 1 can still mark these as taught.', ''),
 ('grade_is_suggested rule: true only when one grade or course was chosen for a standard the source gives to a band (MS and HS NGSS placements, and Wisconsin m/h core ideas placed in some but not all grades). A banded standard copied into every grade of its band keeps grade_band (K-2, 3-5, MS, HS) and has grade_is_suggested false (K-2-ETS1, 3-5-ETS1, MS-ETS1, HS-ETS1 in all four courses, Wisconsin band-level practices and crosscutting concepts). Middle and high school bands are labelled MS and HS.', ''),
 ('Wisconsin Standards for Science rows: K-5 codes carry their own grade (e.g. SCI.LS1.A.4 = Grade 4); K-2 and 3-5 band codes (practices, crosscutting concepts, engineering) appear at each grade of the band; middle (m) and high (h) disciplinary core ideas appear in the grades where their crosswalked NGSS PEs were placed; m and h practices and crosscutting concepts appear in every grade of the band.', ''),
 ('', ''),
 ('Crosswalk', 'h'),
 ('Each row lists related codes as FRAMEWORK:code (NGSS:..., AERO-SCI:..., WI-SCI:...). NGSS to WI-SCI links come from the disciplinary core idea components named in each PE foundation box (NGSS AllDCI.pdf), matched to the Wisconsin DCI code at the same grade or band.', ''),
 ('', ''),
 ('What is DRAFT and needs committee review', 'h'),
 ('1. All "I can" descriptors on the G6-G12 tabs, the WI-SCI tab, and every EE Performance Level (ee_at_target) are DRAFTS written for this project (descriptor_source "draft"): 130 NGSS MS/HS PEs, 253 Wisconsin codes, 43 EEs.', ''),
 ('2. Banded engineering PEs: the school wrote descriptors for K-2-ETS1 only in KG and for 3-5-ETS1 in G4 and G5. Grades 1 and 2 reuse the KG descriptors and Grade 3 reuses the G4 descriptors (marked "school").', ''),
 ('3. School K-5 sheet issues (copied unchanged, please review):', ''),
 ("   a. Code labels: the school sheet writes K-2-ETS1 codes as K-ETS1-1/2/3 and 3-5-ETS1 codes in G5 as 5-ETS1-1/2/3; these were mapped to K-2-ETS1-x and 3-5-ETS1-x. Several codes have spacing/typing variants (e.g. '1 PS4 -1', '2ls2-2', 'PS1-2' in G2, 'PS2-4' in G3, '5 PS3- 1').", ''),
 ("   b. Standard text in the school sheet does not match the code: KG 'K-ETS1-3' shows the K-2-ETS1-1 text; G3 '3-PS2-2' shows text close to 3-PS2-1; G3 'PS2-4' shows the 5-PS1-1 text (particles of matter), and its descriptors are about particles, not magnets.", ''),
 ("   c. All 81 school descriptor rows were reviewed. 23 PEs have school descriptors that describe a different standard: KG K-PS2-2, K-PS3-1, K-ESS2-2, K-2-ETS1-2, K-2-ETS1-3 (the ETS1 rows carry descriptors of the other ETS1 standards); G1 1-LS1-2, 1-PS4-1, 1-PS4-2, 1-PS4-3, 1-PS4-4, 1-ESS1-2 (weather instead of daylight); all 11 G2 PEs (e.g. community helpers, family roles, local places); G3 3-PS2-4 (particles instead of magnets). The school text is kept unchanged in 'descriptors'; each of these rows (NGSS and AERO-SCI, 54 rows including Grades 1-2 reuse of the K-2-ETS1 descriptors) has descriptor_flag 'School descriptor may not match this standard' and a DRAFT 'descriptors_suggested' written for the real standard. In this workbook the school tabs are unchanged; the flags are shaded light yellow on the 'K-5 NGSS-AERO Crosswalk' tab and the suggested descriptor is in its Notes column.", ''),
 ('   d. G4 rows 22-86 hold generic Grade 4 descriptors by DCI (not per standard); they were not imported.', ''),
 ('4. Wisconsin source typos corrected (the code printed in the DPI document is kept in the "source" field): SCI.LS4.A.3 printed in the LS4.C Adaptation row is stored as SCI.LS4.C.3; SCI.LS1.D.2 printed in the Biodiversity and Humans (LS4.D) row is stored as SCI.LS4.D.2; SCI.PS4.C.m printed in the Wave Properties row is stored as SCI.PS4.A.m. The document prints one code as SCI.ESS3.B.3,4 (Grades 3 and 4); it is kept as printed.', ''),
 ('5. Wisconsin "sample performance indicators" are the NGSS PEs (Wisconsin rewords 7 of them slightly: 5-ESS2-1, 5-ESS2-2, MS-ESS2-1, MS-PS3-1, MS-PS3-2, HS-PS2-2, HS-PS2-4, and does not list 3-5-ETS1-3 or MS-LS1-8; the NGSS wording is used). They are not duplicated as WI-SCI rows; only the 12 Wisconsin-only ETS3 indicators (K-ETS3-1 ... HS-ETS3-3) are WI-SCI rows.', ''),
 ('6. PEs with no Essential Element have empty ee_code / ee_text (the Wisconsin EE for Science covers 43 PEs only).', ''),
 ('', ''),
 ('Nothing in this workbook was filled from memory: all standard text comes from the NGSS, AERO, Wisconsin DPI or school documents above.', ''),
]
for i, (t, kind) in enumerate(lines, 1):
    c = ab.cell(i, 1, t)
    c.font = Font(name='Calibri', size=16 if kind == 'title' else 12, bold=kind in ('title', 'h'))
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if kind == 'h': c.fill = fill(BLUE)

# ---------------------------------------------------------------- K-5 crosswalk tab
cw = wb.create_sheet('K-5 NGSS-AERO Crosswalk', 7)
cols = [('Grade', BLUE), ('NGSS Code', BLUE), ('NGSS Performance Expectation', BLUE), ('AERO Code', BLUE), ('AERO Science text', BLUE),
        ('Same wording?', LBLUE), ('Wisconsin Standards for Science (crosswalk)', LBLUE), ('Essential Element', GREEN),
        ('School report card (tab / code as printed)', GREEN), ('In K-5 Science Matrix Map (Science Master Standards.xlsx)', LBLUE), ('Review flag', LBLUE),
        ('School Performance Level used (copied from school tab)', GREEN), ('Notes: suggested descriptor for the real standard (DRAFT)', GREEN)]
header(cw, cols, [8, 14, 60, 14, 60, 12, 34, 16, 30, 22, 55, 80, 80])
YELLOW = fill('FFF2CC')
MATRIX = set('K-LS1-1 K-ESS3-1 K-PS2-1 K-PS3-1 K-ESS2-1 K-ESS3-3 1-LS1-1 1-LS1-2 1-LS3-1 1-PS4-1 1-PS4-2 1-ESS1-1 1-PS4-4 2-LS2-1 2-LS2-2 2-LS4-1 2-PS1-1 2-PS1-4 2-ESS1-1 2-ESS2-1 3-LS1-1 3-LS2-1 3-LS4-3 3-PS2-1 3-ESS2-1 3-ESS3-1 4-LS1-1 4-LS1-2 4-PS3-1 4-PS4-1 4-ESS2-2 4-ESS3-1 5-LS1-1 5-LS2-1 5-PS1-1 5-PS2-1 5-PS3-1 5-ESS1-2 5-ESS3-1'.split())
MATRIX_G = {('K', 'K-2-ETS1-1'), ('4', '3-5-ETS1-1'), ('5', '3-5-ETS1-2')}
FLAG_DESC = {('K', 'K-PS3-1'), ('K', 'K-ESS2-2'), ('1', '1-LS1-2'), ('1', '1-PS4-1'), ('1', '1-PS4-2'), ('1', '1-PS4-3'), ('1', '1-PS4-4'), ('1', '1-ESS1-2'),
             ('2', '2-LS2-1'), ('2', '2-LS2-2'), ('2', '2-LS4-1'), ('2', '2-PS1-1'), ('2', '2-PS1-2'), ('2', '2-PS1-4'), ('2', '2-ESS1-1'), ('2', '2-ESS2-1'), ('2', '2-ESS2-2'), ('2', '2-ESS2-3')}
FLAG_TEXT = {('K', 'K-2-ETS1-3'): 'School row labelled K-ETS1-3 shows the K-2-ETS1-1 text.',
             ('3', '3-PS2-2'): 'School text is close to 3-PS2-1, not 3-PS2-2.',
             ('3', '3-PS2-4'): "School row 'PS2-4' shows the 5-PS1-1 text (particles) and particle descriptors, not the magnet design standard."}
sch = {(x['grade'], x['code']): x for x in school}
def printed(raw):
    m = re.match(r'^(.*?[A-Za-z]{2,3}\s*\d\s*[-\u2010]\s*\d)', raw)
    return m.group(1).strip() if m else raw[:12]
ng_rows = [x for x in rows if x['framework'] == 'NGSS' and x['division'] == 'Elementary']
aero = {(x['grade'], x['code']): x for x in rows if x['framework'] == 'AERO-SCI'}
GO = ['K', '1', '2', '3', '4', '5']
ng_rows.sort(key=lambda x: (GO.index(x['grade']), DISC_ORDER.index(disc(x['code'])), num(x['code'])))
r = 2
for x in ng_rows:
    a = aero[(x['grade'], x['code'])]
    s = sch.get((x['grade'], x['code']))
    same = 'Yes' if re.sub(r'\W', '', a['text'].lower()) == re.sub(r'\W', '', x['text'].lower()) else 'No (AERO wording differs)'
    flags = []
    if x['descriptor_flag']: flags.append(x['descriptor_flag'] + '.')
    if (x['grade'], x['code']) in FLAG_TEXT: flags.append('Review: ' + FLAG_TEXT[(x['grade'], x['code'])])
    if not s: flags.append(f"Not on the school {x['grade'] if x['grade'] != 'K' else 'KG'} report card; descriptors reused from the school row for this banded PE." if x['code'][1:2] == '-' and x['code'][:3] in ('K-2', '3-5') else 'Not on the school report card.')
    inm = 'Yes' if x['code'] in MATRIX or (x['grade'], x['code']) in MATRIX_G else ''
    vals = [x['grade'], x['code'], x['text'], a['code'], a['text'], same,
            '\n'.join(c.split(':', 1)[1] for c in x['crosswalk'] if c.startswith('WI-SCI:')), x['ee_code'],
            (f"{s['tab']} / {printed(s['raw_code'])}" if s else ''),
            inm, ' '.join(flags), perf(x['descriptors']),
            (perf(x['descriptors_suggested']) if x['descriptors_suggested'] else '')]
    for k, v in enumerate(vals, 1):
        c = cw.cell(r, k, v); c.font = Font(name='Calibri', size=11); c.alignment = WRAPTOP if k >= 12 else WRAP; c.border = BORDER
        if x['descriptor_flag'] and k in (11, 12): c.fill = YELLOW
    r += 1

# ---------------------------------------------------------------- MS / HS tabs
grade_tab('G6', '6', 'G6'); grade_tab('G7', '7', 'G7'); grade_tab('G8', '8', 'G8')
grade_tab('Earth Environmental Science', '9', 'Grade 9'); grade_tab('Physics', '10', 'Grade 10')
grade_tab('Biology', '11', 'Grade 11'); grade_tab('Chemistry', '12', 'Grade 12')

# ---------------------------------------------------------------- WI-SCI reference tab
wt = wb.create_sheet('WI-SCI K-12')
cols = [('Strand', BLUE), ('Wisconsin code', BLUE), ('Grade(s)', BLUE), ('Standard text', BLUE), ('Performance Level (DRAFT)', GREEN), ('Crosswalk (NGSS / AERO)', LBLUE)]
header(wt, cols, [26, 20, 14, 80, 90, 30])
byc = OrderedDict()
for x in rows:
    if x['framework'] == 'WI-SCI': byc.setdefault(x['code'], []).append(x)
STR_ORDER = ['Crosscutting Concepts', 'Science and Engineering Practices', 'Life Science', 'Physical Science', 'Earth and Space Science', 'Engineering, Technology, and Applications of Science']
r = 2
for st in STR_ORDER:
    codes = [c for c, xs in byc.items() if xs[0]['strand'] == st]
    if not codes: continue
    banner(wt, r, st, len(cols)); r += 1
    for c in codes:
        xs = byc[c]
        g = ', '.join(OrderedDict.fromkeys(x['grade'] for x in xs))
        xw = list(OrderedDict.fromkeys(cc for x in xs for cc in x['crosswalk']))
        vals = [xs[0]['cluster'], c, g, xs[0]['text'], perf(xs[0]['descriptors']), '\n'.join(xw)]
        for k, v in enumerate(vals, 1):
            cell = wt.cell(r, k, v); cell.font = Font(name='Calibri', size=11, bold=(k == 2)); cell.alignment = WRAPTOP if k in (4, 5) else WRAP; cell.border = BORDER
        r += 1

wb.save(OUT)
print('saved', OUT, wb.sheetnames)
