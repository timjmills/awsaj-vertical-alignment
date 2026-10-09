"""Parse the school's K-5 Science Standards.xlsx report-card tabs into (grade, code) -> descriptors."""
import openpyxl, re, sys, json
wb = openpyxl.load_workbook(sys.argv[1])
GR = {'KG': 'K', 'G1': '1', 'G2': '2', 'G3': '3', 'G4': '4', 'G5': '5'}
LEV = {'emerging': 'cc_emerging', 'approaching': 'cc_approaching', 'at target': 'cc_at_target', 'advanced': 'cc_advanced'}
def normcode(raw, grade):
    s = raw.replace('‐', '-').replace('‑', '-').replace('–', '-').upper()
    s = re.sub(r'\s+', ' ', s).strip()
    m = re.match(r'^(3\s*-\s*5\s*-\s*|K\s*-\s*2\s*-\s*|[K1-5]\s*-?\s*)?(PS|LS|ESS|ETS)\s*(\d)\s*-\s*(\d)', s)
    if not m: return None, None
    pre = (m.group(1) or '').replace(' ', '').rstrip('-')
    if not pre: pre = grade
    code = f"{pre}-{m.group(2)}{m.group(3)}-{m.group(4)}"
    rest = s[m.end():]
    return code, rest
def split_levels(txt):
    out = {}
    if not txt: return out
    t = txt.replace('\r', '')
    pat = re.compile(r'(?im)^\s*(emerging|approaching|at target|advanced)\s*:{0,2}\s*')
    ms = list(pat.finditer(t))
    for i, m in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(t)
        body = t[m.end():end]
        lines = [l.strip() for l in body.split('\n') if l.strip()]
        lines = [re.sub(r'^Descriptor:\s*', '', l) for l in lines]
        out[LEV[m.group(1).lower()]] = ' '.join(lines)
    return out
rows = []
for ws in wb:
    g = GR[ws.title]
    hdr = [str(c.value or '') for c in ws[1]]
    ccol = next(i for i, h in enumerate(hdr) if 'Code & Standard' in h) + 1
    pcol = next(i for i, h in enumerate(hdr) if 'Performance Level' in h) + 1
    domain = ''; concept = ''
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value; b = ws.cell(r, ccol).value; p = ws.cell(r, pcol).value
        if a and not b and not p:
            if str(a).strip() in ('Life Science', 'Physical Science', 'Earth Science', 'Engineering Design', 'Earth and Space Science'):
                domain = str(a).strip(); continue
            if g == '4' and r > 21: break
        if a: concept = ' '.join(str(a).split())
        if (not b or not str(b).strip()) and ccol > 2 and ws.cell(r, 2).value and p: b = ws.cell(r, 2).value
        if not b: continue
        code, rest = normcode(str(b), g)
        if g == '5' and code and re.match(r'^5-ETS1', code): code = '3-5-' + code[2:]
        if g == 'K' and code and code.startswith('K-ETS1'): code = 'K-2-' + code[2:]
        if g in ('4',) and code and code.startswith('3-5-ETS'): pass
        lv = split_levels(str(p or ''))
        rows.append({'tab': ws.title, 'row': r, 'grade': g, 'domain': domain or ws.cell(2, 1).value, 'concept': concept,
                     'raw_code': ' '.join(str(b).split())[:40], 'code': code, 'school_text': ' '.join(str(b).split()),
                     'levels': lv, 'n_levels': len(lv)})
json.dump(rows, open(sys.argv[2], 'w'), indent=1, ensure_ascii=False)
for x in rows: print(x['tab'], x['row'], x['code'], '|', x['raw_code'][:30], '| levels', x['n_levels'], '|', x['domain'], '/', x['concept'][:30])
