"""Parse Wisconsin Standards for Science (2017, accessible .docx from dpi.wi.gov) into WI-SCI codes and the embedded
NGSS sample performance indicators (PEs)."""
import docx, re, sys, json
d = docx.Document(sys.argv[1])
CODE = re.compile(r'(SCI\.(?:CC\d|SEP\d(?:\.[AB])?|(?:LS|PS|ESS|ETS)\d\.[A-E])\.(?:K-2|3-5|K|[1-5](?:,[1-5])?|m|h))(?![\w-])')
PE = re.compile(r'((?:K|[1-5]|K-2|3-5|MS|HS)-(?:PS|LS|ESS|ETS)\d-\d+)\.\s*')
codes = {}   # code -> dict
pes = {}
order = 0
fixes = []
# map paragraph headings: find "Standard SCI.XX:" statements for strand names
std_text = {}
for p in d.paragraphs:
    m = re.match(r'\s*Standard (SCI\.[A-Z]+\d?):\s*(.*)', p.text)
    if m: std_text.setdefault(m.group(1), ' '.join(m.group(2).split()))
def clean(s): return ' '.join(s.replace(' ', ' ').split())
for ti, t in enumerate(d.tables):
    rows = t.rows
    hdr = [c.text.strip() for c in rows[0].cells]
    if hdr and hdr[0].startswith('Learning Priority'):
        bands = hdr[1:]
        for r in rows[1:]:
            cells = [c.text for c in r.cells]
            label = clean(cells[0])
            seen = []
            rcells = list(r.cells)
            for bi, cell in enumerate(cells[1:]):
                tc = rcells[bi + 1]._tc
                if any(tc is x for x in seen): continue
                seen.append(tc)
                parts = CODE.split(cell)
                # parts: [pre, code, text, code, text...]
                for k in range(1, len(parts), 2):
                    printed = parts[k]; txt = parts[k + 1] if k + 1 < len(parts) else ''
                    code = printed
                    lm = re.match(r'(?:SCI\.)?((?:LS|PS|ESS|ETS)\d\.[A-E]|CC\d|SEP\d(?:\.[AB])?)\s*:\s*(.*)', label)
                    rowcomp = lm.group(1) if lm else None
                    if lm and lm.group(2).startswith('Biodiversity and Humans'): rowcomp = 'LS4.D'  # source labels this row SCI.LS1.D
                    pc = printed.split('.', 1)[1].rsplit('.', 1)[0]
                    if rowcomp and pc != rowcomp and not pc.startswith(('CC', 'SEP')):
                        code = 'SCI.' + rowcomp + '.' + printed.rsplit('.', 1)[1]
                        fixes.append((printed, code, label))
                    lines = [clean(x) for x in txt.split('\n') if clean(x)]
                    body = ' '.join(lines)
                    lab2 = re.sub(r"\s*\(cont.d\)", '', label)
                    if rowcomp == 'LS4.D': lab2 = 'SCI.LS4.D: Biodiversity and Humans'
                    e = codes.setdefault(code, {'code': code, 'printed_code': printed, 'label': lab2, 'band_col': bands[bi], 'parts': [], 'order': order})
                    order += 1
                    if body: e['parts'].append(body)
    elif hdr and hdr[0].startswith('Grades'):
        for r in rows:
            band = r.cells[0].text.strip()
            txt = r.cells[1].text
            ms = list(PE.finditer(txt))
            for i, m in enumerate(ms):
                end = ms[i + 1].start() if i + 1 < len(ms) else len(txt)
                pes.setdefault(m.group(1), {'code': m.group(1), 'text': clean(txt[m.end():end]), 'band': band})
for c in codes.values():
    c['text'] = ' '.join(c['parts']).strip()
    del c['parts']
json.dump({'fixes': fixes, 'codes': codes, 'pes': pes, 'standards': std_text}, open(sys.argv[2], 'w'), indent=1, ensure_ascii=False)
from collections import Counter
print('WI codes', len(codes), Counter(c.rsplit('.', 1)[1] for c in codes))
print('PEs', len(pes), Counter(p.split('-')[0] for p in pes))
print('standards', std_text.keys())
print([c for c in codes if not codes[c]['text']]); print('fixes', fixes)
